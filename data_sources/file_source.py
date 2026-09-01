"""Normalized file-based product data sources for CSV and Excel imports."""

from __future__ import annotations

import csv
from datetime import datetime, timezone
import hashlib
from pathlib import Path
import re
from typing import Any, Iterable

from data_sources.base import BaseDataSource, DataSourceRecord, DataSourceResult, DataSourceValidationError
from scrapers.models import ScrapedProductData
from scrapers.product_page_scraper import ProductPageScraper


class FileDataSource(BaseDataSource):
	"""Load normalized product rows from CSV or XLSX files."""

	SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}
	SUPPORTED_COLUMNS = (
		"platform",
		"product_name",
		"brand",
		"model",
		"current_price",
		"old_price",
		"discount_percentage",
		"currency",
		"seller_name",
		"seller_rating",
		"availability",
		"product_url",
		"external_product_id",
		"visible_sales_count",
	)

	def load(self, source: str | Path) -> DataSourceResult:
		path = Path(source)
		suffix = path.suffix.lower()
		if suffix not in self.SUPPORTED_EXTENSIONS:
			return DataSourceResult(
				source_failure_reason="unsupported_file_type",
				source_failure_details="Only .csv and .xlsx files are supported for file imports.",
			)

		try:
			if suffix == ".csv":
				rows = self._read_csv_rows(path)
			else:
				rows = self._read_excel_rows(path)
		except (OSError, ValueError) as error:
			return DataSourceResult(
				source_failure_reason="file_read_failed",
				source_failure_details=str(error),
			)

		records: list[DataSourceRecord] = []
		validation_errors: list[DataSourceValidationError] = []
		for row_number, row in rows:
			normalized_data, row_errors = self._normalize_row(row)
			if row_errors:
				validation_errors.append(
					DataSourceValidationError(row_number=row_number, messages=row_errors)
				)
				continue

			records.append(DataSourceRecord(data=normalized_data, row_number=row_number))

		return DataSourceResult(records=records, validation_errors=validation_errors)

	def _read_csv_rows(self, path: Path) -> list[tuple[int, dict[str, Any]]]:
		with path.open("r", encoding="utf-8-sig", newline="") as handle:
			reader = csv.DictReader(handle)
			if reader.fieldnames is None:
				raise ValueError("The CSV file does not contain a header row.")

			rows: list[tuple[int, dict[str, Any]]] = []
			for row_number, raw_row in enumerate(reader, start=2):
				rows.append((row_number, self._normalize_input_mapping(raw_row)))

		return rows

	def _read_excel_rows(self, path: Path) -> list[tuple[int, dict[str, Any]]]:
		try:
			from openpyxl import load_workbook
			from openpyxl.utils.exceptions import InvalidFileException
		except ImportError as error:
			raise ValueError("openpyxl is required for .xlsx imports.") from error

		try:
			workbook = load_workbook(path, read_only=True, data_only=True)
		except (OSError, InvalidFileException, ValueError) as error:
			raise ValueError(f"Failed to read Excel file: {error}") from error

		try:
			worksheet = workbook.active
			row_iterator = worksheet.iter_rows(values_only=True)
			headers = next(row_iterator, None)
			if headers is None:
				raise ValueError("The Excel file does not contain a header row.")

			normalized_headers = [self._normalize_header(header) for header in headers]
			rows: list[tuple[int, dict[str, Any]]] = []
			for row_number, values in enumerate(row_iterator, start=2):
				row: dict[str, Any] = {}
				for index, header in enumerate(normalized_headers):
					if not header:
						continue
					row[header] = values[index] if index < len(values) else None
				rows.append((row_number, row))
			return rows
		finally:
			workbook.close()

	def _normalize_row(self, row: dict[str, Any]) -> tuple[ScrapedProductData | None, list[str]]:
		platform = self._normalize_platform(row.get("platform"))
		product_name = self._clean_text(row.get("product_name"))
		brand = self._clean_text(row.get("brand"))
		model = self._clean_text(row.get("model"))
		seller_name = self._clean_text(row.get("seller_name"))
		currency = self._clean_text(row.get("currency"))
		availability = self._clean_text(row.get("availability"))
		product_url = self._clean_text(row.get("product_url"))
		external_product_id = self._clean_text(row.get("external_product_id"))

		errors: list[str] = []
		if platform is None:
			errors.append("platform is missing")
		if product_name is None:
			errors.append("product_name is missing")

		current_price_raw = row.get("current_price")
		if self._is_missing(current_price_raw):
			errors.append("current_price is missing")
			current_price = None
		else:
			current_price = ProductPageScraper.normalize_money(current_price_raw)
			if current_price is None:
				errors.append("invalid current_price format")

		old_price, old_price_error = self._parse_optional_money(row.get("old_price"), "old_price")
		if old_price_error is not None:
			errors.append(old_price_error)

		discount_percentage, discount_error = self._parse_optional_decimal(
			row.get("discount_percentage"),
			"discount_percentage",
		)
		if discount_error is not None:
			errors.append(discount_error)

		seller_rating, seller_rating_error = self._parse_optional_decimal(
			row.get("seller_rating"),
			"seller_rating",
		)
		if seller_rating_error is not None:
			errors.append(seller_rating_error)

		visible_sales_count, visible_sales_count_error = self._parse_optional_int(
			row.get("visible_sales_count"),
			"visible_sales_count",
		)
		if visible_sales_count_error is not None:
			errors.append(visible_sales_count_error)

		if errors:
			return None, errors

		if external_product_id is None:
			external_product_id = self._generate_external_product_id(
				platform=platform,
				brand=brand,
				model=model,
				product_name=product_name,
				seller_name=seller_name,
			)

		if product_url is None:
			product_url = self._build_import_product_url(platform, external_product_id)

		return (
			ScrapedProductData(
				platform=platform,
				product_name=product_name,
				brand=brand,
				model=model,
				current_price=current_price,
				old_price=old_price,
				discount_percentage=discount_percentage,
				currency=currency.upper() if currency is not None else None,
				seller_name=seller_name,
				seller_rating=seller_rating,
				availability=availability,
				product_url=product_url,
				external_product_id=external_product_id,
				scraped_at=datetime.now(timezone.utc),
				visible_sales_count=visible_sales_count,
			),
			errors,
		)

	@classmethod
	def _normalize_input_mapping(cls, row: dict[str, Any]) -> dict[str, Any]:
		normalized: dict[str, Any] = {}
		for key, value in row.items():
			normalized_key = cls._normalize_header(key)
			if not normalized_key:
				continue
			normalized[normalized_key] = value
		return normalized

	@staticmethod
	def _normalize_header(value: Any) -> str:
		if value is None:
			return ""
		return str(value).strip().lower()

	@staticmethod
	def _normalize_platform(value: Any) -> str | None:
		cleaned = FileDataSource._clean_text(value)
		if cleaned is None:
			return None
		return cleaned.lower()

	@staticmethod
	def _clean_text(value: Any) -> str | None:
		if value is None:
			return None
		cleaned = str(value).strip()
		if not cleaned:
			return None
		return cleaned

	@staticmethod
	def _is_missing(value: Any) -> bool:
		if value is None:
			return True
		if isinstance(value, str):
			return not value.strip()
		return False

	@staticmethod
	def _parse_optional_money(value: Any, field_name: str) -> tuple[Any, str | None]:
		if FileDataSource._is_missing(value):
			return None, None
		parsed = ProductPageScraper.normalize_money(value)
		if parsed is None:
			return None, f"invalid {field_name} format"
		return parsed, None

	@staticmethod
	def _parse_optional_decimal(value: Any, field_name: str) -> tuple[Any, str | None]:
		if FileDataSource._is_missing(value):
			return None, None
		parsed = ProductPageScraper.normalize_decimal(str(value))
		if parsed is None:
			return None, f"invalid {field_name} format"
		return parsed, None

	@staticmethod
	def _parse_optional_int(value: Any, field_name: str) -> tuple[int | None, str | None]:
		if FileDataSource._is_missing(value):
			return None, None
		if isinstance(value, bool):
			return None, f"invalid {field_name} format"
		if isinstance(value, int):
			return value, None
		if isinstance(value, float):
			if value.is_integer():
				return int(value), None
			return None, f"invalid {field_name} format"
		cleaned = str(value).strip()
		if cleaned.isdigit():
			return int(cleaned), None
		return None, f"invalid {field_name} format"

	@staticmethod
	def _generate_external_product_id(
		platform: str,
		brand: str | None,
		model: str | None,
		product_name: str,
		seller_name: str | None,
	) -> str:
		identity_parts = (
			platform.strip().lower(),
			(brand or "").strip().lower(),
			(model or product_name).strip().lower(),
			(seller_name or "").strip().lower(),
		)
		digest = hashlib.sha256("|".join(identity_parts).encode("utf-8")).hexdigest()[:16].upper()
		return f"import-generated-{digest}"

	@staticmethod
	def _build_import_product_url(platform: str, external_product_id: str) -> str:
		platform_slug = re.sub(r"[^a-z0-9._-]+", "-", platform.lower()).strip("-") or "unknown"
		return f"import://{platform_slug}/{external_product_id}"